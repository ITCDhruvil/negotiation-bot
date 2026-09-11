"""Turn an RFQ document (or pasted text) into a filled draft.

Heuristic parse always runs so tests and labeled PDFs work without spending
tokens. When a live OpenAI key is present, the chat model overlays any fields
the regex pass missed. Deal-engine numbers are never invented here beyond
leaving blanks for the buyer to confirm.
"""

from __future__ import annotations

import base64
import json
import logging
import os
import re
from typing import Any

from app.models import AlternateQuoteDraft, RfqDraft, RfqPartDraft

logger = logging.getLogger("aria.rfq")

MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 24_000

_TEXT_TYPES = {
    "text/plain",
    "text/markdown",
    "text/csv",
    "application/json",
    "application/xml",
    "text/xml",
}
_IMAGE_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif"}
_PDF_TYPES = {"application/pdf"}
_DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}

_REQUIRED_PART = ("part_name", "vendor_name", "quantity", "vendor_quoted_unit_price")


def _norm_space(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text).strip()


def _to_int(raw: str | None) -> int | None:
    if raw is None:
        return None
    cleaned = raw.strip().lower().replace("₹", "").replace("rs.", "").replace("rs", "").replace("inr", "")
    cleaned = cleaned.replace(",", "").replace(" ", "")
    match = re.search(r"-?\d+", cleaned)
    if not match:
        return None
    try:
        return int(match.group(0))
    except ValueError:
        return None


def _label_value(text: str, labels: str) -> str | None:
    pattern = rf"(?im)^(?:{labels})\s*[:\-]\s*(.+)$"
    match = re.search(pattern, text)
    if not match:
        return None
    value = _norm_space(match.group(1))
    return value or None


def _first(*values: str | None) -> str | None:
    for value in values:
        if value:
            return value
    return None


def _lead_days(raw: str | None) -> int | None:
    if not raw:
        return None
    weeks = re.search(r"(\d+)\s*week", raw, re.I)
    if weeks:
        return int(weeks.group(1)) * 7
    return _to_int(raw)


def _split_part_blocks(text: str) -> list[str]:
    chunks = re.split(r"(?im)^\s*(?:part\s*(?:\d+|name)|line\s*item)\b", text)
    if len(chunks) <= 1:
        return [text]
    blocks: list[str] = []
    for chunk in chunks[1:]:
        body = chunk.strip()
        if not body:
            continue
        if not re.match(r"(?i)^\s*(?:name)?\s*[:\-]", body):
            body = "Part name: " + body
        else:
            body = "Part " + body
        blocks.append(body)
    return blocks or [text]


def _parse_alternates(text: str) -> list[AlternateQuoteDraft]:
    found: list[AlternateQuoteDraft] = []
    for match in re.finditer(
        r"(?im)^(?:alternate(?:\s*vendor)?|alt(?:ernate)?\s*quote)\s*[:\-]\s*(.+?)(?:\s+at\s+|[:,]\s*)₹?\s*([\d,]+)",
        text,
    ):
        name = _norm_space(re.sub(r"\bat\b.*$", "", match.group(1), flags=re.I))
        price = _to_int(match.group(2))
        if name and price:
            found.append(AlternateQuoteDraft(vendor_name=name, unit_price=price))
    return found


def _parse_part(block: str) -> RfqPartDraft:
    part = RfqPartDraft(
        part_name=_first(
            _label_value(block, r"part(?:\s*name)?|item(?:\s*name)?|component"),
        )
        or "",
        category=_label_value(block, r"category|commodity|group") or "",
        vendor_name=_first(
            _label_value(block, r"vendor(?:\s*name)?|supplier(?:\s*name)?"),
        )
        or "",
        vendor_id=_label_value(block, r"vendor\s*id|supplier\s*id") or "",
        vendor_quoted_unit_price=_to_int(
            _first(
                _label_value(block, r"(?:vendor\s*)?(?:quoted\s*)?(?:unit\s*)?price|quote|rate"),
            )
        ),
        quantity=_to_int(_label_value(block, r"quantity|qty|units?")),
        moq=_to_int(_label_value(block, r"moq|minimum\s*order(?:\s*qty)?")),
        target_unit_price=_to_int(_label_value(block, r"target(?:\s*unit)?(?:\s*price)?")),
        max_acceptable_unit_price=_to_int(
            _label_value(block, r"max(?:imum)?(?:\s*acceptable)?(?:\s*unit)?(?:\s*price)?|walk[- ]?away")
        ),
        lead_time_days=_lead_days(_label_value(block, r"lead\s*time(?:\s*\(?days\)?)?")),
        payment_terms_default=_label_value(block, r"payment\s*terms(?:\s*default)?") or "",
        spec_blurb=_first(
            _label_value(block, r"spec(?:ification)?s?|blurb|description|notes"),
        )
        or "",
        target_lead_time_days=_lead_days(_label_value(block, r"target\s*lead\s*time")),
        max_acceptable_lead_time_days=_lead_days(
            _label_value(block, r"max(?:imum)?(?:\s*acceptable)?\s*lead\s*time")
        ),
        max_acceptable_moq=_to_int(_label_value(block, r"max(?:imum)?(?:\s*acceptable)?\s*moq")),
        preferred_payment_terms=_label_value(block, r"preferred\s*payment\s*terms") or "",
        fastest_payment_terms=_label_value(block, r"fastest\s*payment\s*terms") or "",
        min_warranty_months=_to_int(_label_value(block, r"(?:min(?:imum)?\s*)?warranty(?:\s*months)?")),
        alternate_vendor_quotes=_parse_alternates(block),
    )
    return annotate_missing(part)


def annotate_missing(part: RfqPartDraft) -> RfqPartDraft:
    missing = [name for name in _REQUIRED_PART if not getattr(part, name)]
    part.missing = missing
    return part


def annotate_draft(draft: RfqDraft) -> RfqDraft:
    warnings = list(draft.warnings)
    if not draft.title:
        warnings.append("RFQ title is missing.")
    if not draft.plant:
        warnings.append("Plant is missing.")
    if not draft.parts:
        warnings.append("No part lines were found.")
        draft.parts = [annotate_missing(RfqPartDraft())]
    for part in draft.parts:
        annotate_missing(part)
        quote = part.vendor_quoted_unit_price or 0
        qty = part.quantity or 0
        if quote and qty and quote * qty > 1_000_000:
            warnings.append(
                f"{part.part_name or 'This line'} quotes ₹{quote * qty:,} total — above Aria's ₹10L automated ceiling."
            )
    # unique
    seen: set[str] = set()
    draft.warnings = [item for item in warnings if not (item in seen or seen.add(item))]
    return draft


def parse_text(text: str) -> RfqDraft:
    cleaned = text.replace("\r\n", "\n").strip()
    header = cleaned.split("\nPart", 1)[0]
    title = _first(
        _label_value(header, r"rfq\s*title|request\s*title|title|subject"),
        _label_value(cleaned, r"rfq\s*title|request\s*title|title|subject"),
    )
    draft = RfqDraft(
        request_id=_label_value(cleaned, r"rfq(?:\s*id)?|request\s*id") or "",
        title=title or "",
        plant=_label_value(cleaned, r"plant|location|site") or "",
        needed_by=_label_value(cleaned, r"needed\s*by|due(?:\s*date)?|delivery\s*date"),
        parts=[_parse_part(block) for block in _split_part_blocks(cleaned)],
        source="heuristic",
    )
    if len(draft.parts) == 1 and not draft.parts[0].part_name:
        # Single unlabeled blob: still try whole-document labels.
        draft.parts = [_parse_part(cleaned)]
    return annotate_draft(draft)


def extract_text_from_bytes(filename: str, content: bytes, content_type: str | None) -> tuple[str, str]:
    """Return (text, kind) where kind is text|pdf|docx|image|empty."""
    name = (filename or "").lower()
    ctype = (content_type or "").split(";")[0].strip().lower()
    if ctype in _IMAGE_TYPES or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        return "", "image"
    if ctype in _PDF_TYPES or name.endswith(".pdf"):
        return _pdf_text(content), "pdf"
    if ctype in _DOCX_TYPES or name.endswith((".docx", ".doc")):
        return _docx_text(content), "docx"
    if ctype in _TEXT_TYPES or name.endswith((".txt", ".md", ".csv", ".json")):
        return content.decode("utf-8", errors="replace"), "text"
    # Fall through: try UTF-8, then PDF.
    try:
        decoded = content.decode("utf-8")
        if decoded.strip():
            return decoded, "text"
    except Exception:
        pass
    pdf_try = _pdf_text(content)
    if pdf_try.strip():
        return pdf_try, "pdf"
    return "", "empty"


def _pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
        import io

        reader = PdfReader(io.BytesIO(content))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    except Exception:
        logger.exception("PDF text extract failed")
        return ""


def _docx_text(content: bytes) -> str:
    try:
        import io
        from docx import Document

        doc = Document(io.BytesIO(content))
        lines = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except Exception:
        logger.exception("DOCX text extract failed")
        return ""


def _llm_allowed() -> bool:
    if os.environ.get("ARIA_ALLOW_MOCK") == "1":
        return False
    from app.config import get_settings

    return bool(get_settings().openai_api_key)


async def enrich_with_llm(
    draft: RfqDraft,
    *,
    text: str,
    image: bytes | None = None,
    mime: str | None = None,
) -> RfqDraft:
    if not _llm_allowed():
        return draft
    from openai import AsyncOpenAI

    from app.config import get_settings

    settings = get_settings()
    client = AsyncOpenAI(api_key=settings.openai_api_key)
    model = settings.openai_model or "gpt-4o"
    schema_hint = {
        "request_id": "string or null",
        "title": "string or null",
        "plant": "string or null",
        "needed_by": "YYYY-MM-DD or null",
        "parts": [
            {
                "part_name": "string",
                "category": "string or null",
                "vendor_name": "string or null",
                "vendor_id": "string or null",
                "vendor_quoted_unit_price": "integer INR or null",
                "quantity": "integer or null",
                "moq": "integer or null",
                "target_unit_price": "integer INR or null",
                "max_acceptable_unit_price": "integer INR or null",
                "lead_time_days": "integer or null",
                "payment_terms_default": "string or null",
                "spec_blurb": "string or null",
                "target_lead_time_days": "integer or null",
                "max_acceptable_lead_time_days": "integer or null",
                "max_acceptable_moq": "integer or null",
                "preferred_payment_terms": "string or null",
                "fastest_payment_terms": "string or null",
                "min_warranty_months": "integer or null",
                "alternate_vendor_quotes": [{"vendor_name": "string", "unit_price": "integer"}],
            }
        ],
    }
    system = (
        "You extract a SKODA Auto India buyer-side parts RFQ into JSON. "
        "Use only facts present in the document. Never invent a vendor, price, or quantity. "
        "Prices are integers in INR with no commas or currency symbols. "
        "If a field is absent, use null. Return a single JSON object matching the schema."
    )
    user_text = (
        "Schema:\n"
        + json.dumps(schema_hint)
        + "\n\nDocument:\n"
        + (text or "[image attached — read the RFQ from the image]")[:MAX_TEXT_CHARS]
    )
    content: Any
    if image:
        b64 = base64.b64encode(image).decode("ascii")
        media = mime or "image/png"
        content = [
            {"type": "text", "text": user_text},
            {"type": "image_url", "image_url": {"url": f"data:{media};base64,{b64}"}},
        ]
    else:
        content = user_text
    try:
        response = await client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": content},
            ],
            response_format={"type": "json_object"},
            max_tokens=1800,
        )
    except Exception:
        logger.exception("RFQ LLM extract failed")
        return draft
    raw = (response.choices[0].message.content or "").strip()
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError:
        logger.warning("RFQ LLM extract returned non-JSON")
        return draft
    merged = _merge_llm(draft, payload)
    merged.source = "llm"
    return annotate_draft(merged)


def _as_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    return _to_int(str(value))


def _as_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _fill(current: str | int | None, incoming: Any, *, numeric: bool = False):
    if numeric:
        parsed = _as_int(incoming)
        return current if current is not None else parsed
    text = _as_str(incoming)
    return current if current else text


def _merge_llm(draft: RfqDraft, payload: dict) -> RfqDraft:
    draft.title = _fill(draft.title, payload.get("title"))
    draft.plant = _fill(draft.plant, payload.get("plant"))
    draft.request_id = _fill(draft.request_id, payload.get("request_id"))
    if not draft.needed_by:
        needed = payload.get("needed_by")
        draft.needed_by = _as_str(needed) or None
    llm_parts = payload.get("parts") if isinstance(payload.get("parts"), list) else []
    if not llm_parts:
        return draft
    if len(draft.parts) == 1 and not draft.parts[0].part_name and llm_parts:
        draft.parts = [_part_from_llm(item) for item in llm_parts]
        return draft
    for index, item in enumerate(llm_parts):
        if not isinstance(item, dict):
            continue
        if index < len(draft.parts):
            draft.parts[index] = _overlay_part(draft.parts[index], item)
        else:
            draft.parts.append(_part_from_llm(item))
    return draft


def _part_from_llm(item: dict) -> RfqPartDraft:
    return _overlay_part(RfqPartDraft(), item)


def _overlay_part(part: RfqPartDraft, item: dict) -> RfqPartDraft:
    part.part_name = _fill(part.part_name, item.get("part_name"))
    part.category = _fill(part.category, item.get("category"))
    part.vendor_name = _fill(part.vendor_name, item.get("vendor_name"))
    part.vendor_id = _fill(part.vendor_id, item.get("vendor_id"))
    part.vendor_quoted_unit_price = _fill(
        part.vendor_quoted_unit_price, item.get("vendor_quoted_unit_price"), numeric=True
    )
    part.quantity = _fill(part.quantity, item.get("quantity"), numeric=True)
    part.moq = _fill(part.moq, item.get("moq"), numeric=True)
    part.target_unit_price = _fill(part.target_unit_price, item.get("target_unit_price"), numeric=True)
    part.max_acceptable_unit_price = _fill(
        part.max_acceptable_unit_price, item.get("max_acceptable_unit_price"), numeric=True
    )
    part.lead_time_days = _fill(part.lead_time_days, item.get("lead_time_days"), numeric=True)
    part.payment_terms_default = _fill(part.payment_terms_default, item.get("payment_terms_default"))
    part.spec_blurb = _fill(part.spec_blurb, item.get("spec_blurb"))
    part.target_lead_time_days = _fill(
        part.target_lead_time_days, item.get("target_lead_time_days"), numeric=True
    )
    part.max_acceptable_lead_time_days = _fill(
        part.max_acceptable_lead_time_days, item.get("max_acceptable_lead_time_days"), numeric=True
    )
    part.max_acceptable_moq = _fill(part.max_acceptable_moq, item.get("max_acceptable_moq"), numeric=True)
    part.preferred_payment_terms = _fill(part.preferred_payment_terms, item.get("preferred_payment_terms"))
    part.fastest_payment_terms = _fill(part.fastest_payment_terms, item.get("fastest_payment_terms"))
    part.min_warranty_months = _fill(part.min_warranty_months, item.get("min_warranty_months"), numeric=True)
    alts = item.get("alternate_vendor_quotes")
    if isinstance(alts, list) and not part.alternate_vendor_quotes:
        for alt in alts:
            if not isinstance(alt, dict):
                continue
            name = _as_str(alt.get("vendor_name"))
            price = _as_int(alt.get("unit_price"))
            if name and price:
                part.alternate_vendor_quotes.append(AlternateQuoteDraft(vendor_name=name, unit_price=price))
    return annotate_missing(part)


async def parse_document(
    *,
    filename: str,
    content: bytes,
    content_type: str | None,
    pasted_text: str = "",
) -> RfqDraft:
    if len(content) > MAX_BYTES:
        raise ValueError("File is larger than 8 MB.")
    text, kind = extract_text_from_bytes(filename, content, content_type) if content else ("", "empty")
    if pasted_text.strip():
        text = (text + "\n" + pasted_text).strip() if text.strip() else pasted_text.strip()
        if kind == "empty":
            kind = "text"
    if not text.strip() and kind != "image":
        draft = RfqDraft(
            filename=filename or None,
            source="empty",
            warnings=["Could not read text from that file. Paste the RFQ or fill the form."],
            parts=[annotate_missing(RfqPartDraft())],
        )
        return draft
    draft = parse_text(text) if text.strip() else RfqDraft(parts=[annotate_missing(RfqPartDraft())])
    draft.filename = filename or None
    image = content if kind == "image" else None
    mime = content_type if kind == "image" else None
    if kind == "image" or (kind == "pdf" and len(text.strip()) < 40):
        draft = await enrich_with_llm(draft, text=text, image=image, mime=mime)
    elif text.strip():
        needs_llm = (not draft.title) or any(part.missing for part in draft.parts)
        if needs_llm:
            draft = await enrich_with_llm(draft, text=text)
        else:
            draft.source = "heuristic"
    return annotate_draft(draft)
